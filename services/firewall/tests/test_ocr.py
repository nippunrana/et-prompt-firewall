import io
import pytest
from fastapi.testclient import TestClient
import docx

from app.main import app
from app.ocr import (
    extract_from_docx_sync,
    format_ocr_boxes_to_markdown,
    html_table_to_markdown,
    process_document,
)

client = TestClient(app)


def test_html_table_to_markdown():
    html = "<table><tr><td>Header 1</td><td>Header 2</td></tr><tr><td>Val 1</td><td>Val 2</td></tr></table>"
    md = html_table_to_markdown(html)
    assert "| Header 1 | Header 2 |" in md
    assert "| Val 1 | Val 2 |" in md


def test_format_ocr_boxes_to_markdown():
    # Synthetic OCR boxes
    # Format: [box, text, score]
    # box: [[x1, y1], [x2, y1], [x2, y2], [x1, y2]]
    sample_boxes = [
        [[[10, 10], [100, 10], [100, 30], [10, 30]], "Invoice #1234", 0.99],
        [[[10, 50], [80, 50], [80, 70], [10, 70]], "Item A", 0.98],
        [[[200, 50], [250, 50], [250, 70], [200, 70]], "$100", 0.97],
    ]
    md = format_ocr_boxes_to_markdown(sample_boxes)
    assert "Invoice #1234" in md
    assert "Item A" in md
    assert "$100" in md


def test_docx_extraction():
    doc = docx.Document()
    doc.add_heading("Candidate Resume", level=1)
    doc.add_paragraph("Experienced Software Engineer.")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Years"
    table.rows[1].cells[0].text = "Python"
    table.rows[1].cells[1].text = "5"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    result = extract_from_docx_sync(docx_bytes)
    assert "Candidate Resume" in result
    assert "Experienced Software Engineer." in result
    assert "| Skill | Years |" in result
    assert "| Python | 5 |" in result


def test_endpoint_docx():
    doc = docx.Document()
    doc.add_paragraph("Test Document Content")
    buf = io.BytesIO()
    doc.save(buf)

    response = client.post(
        "/extract-text",
        files={"file": ("test.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    assert data["method"] == "word_docx"
    assert "Test Document Content" in data["text"]


def test_endpoint_empty_file():
    response = client.post(
        "/extract-text",
        files={"file": ("empty.png", b"", "image/png")},
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()
